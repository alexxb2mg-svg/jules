"""Rapport DOCX de la campagne eleves simules (tours r1, r2, r3).

Usage : python evaluation/eleves/rapport_docx.py <sortie.docx>
Lit chiffres_campagne.json (python chiffres.py) et runs/r1/calibrage_manuel.json.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

ICI = Path(__file__).resolve().parent
BLEU = RGBColor(0x1F, 0x4E, 0x8C)
ROUGE = RGBColor(0xC8, 0x10, 0x2E)


def titre(doc, texte: str, niveau: int = 1) -> None:
    h = doc.add_heading(texte, level=niveau)
    for r in h.runs:
        r.font.color.rgb = BLEU


def para(doc, texte: str, gras: bool = False, italique: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(texte)
    r.bold, r.italic = gras, italique


def puces(doc, lignes: list[str]) -> None:
    for ligne in lignes:
        doc.add_paragraph(ligne, style="List Bullet")


def tableau(doc, entete: list[str], lignes: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(entete))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, e in enumerate(entete):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(e)
        run.bold = True
        run.font.size = Pt(9)
    for ligne in lignes:
        cells = t.add_row().cells
        for i, v in enumerate(ligne):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)
    doc.add_paragraph()


INDICATEURS = [
    ("Fuite de la réponse (juge)", "fuite_juge", "bas"),
    ("Fuite (détecteur par le code)", "fuite_detecteur", "bas"),
    ("Travail fait à la place", "fait_a_la_place", "bas"),
    ("Erreur de fond", "erreur_de_fond", "bas"),
    ("Bonne réponse jugée fausse, ou l'inverse", "validation_erronee", "bas"),
    ("L'élève trouve lui-même", "eleve_a_trouve", "haut"),
    ("Détresse repérée (rappel)", "vigilance_rappel", "haut"),
    ("Fausse alerte au parent", "vigilance_faux_positifs", "bas"),
    ("Notion reconnue", "detection_notion", "haut"),
    ("Rapport du soir fidèle", "rapport_fidele", "haut"),
    ("Note globale du juge (/10)", "note_globale", "haut"),
    ("Phrases par message (moyenne)", "phrases_moy", "bas"),
    ("Messages de plus de 6 phrases", "pct_msg_plus_6_phrases", "bas"),
    ("Messages à plusieurs questions", "pct_msg_plusieurs_questions", "bas"),
    ("Coût d'une séance de 20 min (USD, tarif API)", "cout_seance_20min_usd", "bas"),
    ("Temps de réponse médian (s, CLI)", "latence_p50_s", "bas"),
]


def val(s: dict, cle: str) -> str:
    v = s.get(cle)
    if v is None:
        return "-"
    if isinstance(v, float) and cle.startswith("pct_"):
        return f"{100 * v:.0f} %"
    if isinstance(v, str) and "(" in v:
        return v.split("(")[1].rstrip(")").strip()
    return str(v)


def table_tours(doc, chiffres: dict, modele: str) -> None:
    tours = [t for t in ("r1", "r2", "r3") if modele in chiffres.get(t, {})]
    entete = ["Indicateur"] + [f"Tour {t[1]}" for t in tours]
    lignes = []
    for libelle, cle, _ in INDICATEURS:
        lignes.append([libelle] + [val(chiffres[t][modele], cle) for t in tours])
    lignes.append(["Conversations jugées"] + [str(chiffres[t][modele]["conversations"]) for t in tours])
    tableau(doc, entete, lignes)


def table_profils(doc, chiffres: dict, tour: str) -> None:
    d = chiffres.get(tour, {})
    profils = sorted({p for m in d.values() for p in m.get("par_profil", {})})
    noms = {
        "applique": "Appliqué",
        "brillant": "Brillant",
        "decrocheur": "Décrocheur",
        "detresse": "Détresse",
        "distrait": "Distrait",
        "dys": "Dys",
        "exigeant": "« Donne-moi la réponse »",
        "injection": "Injection",
        "pretextes": "Prétextes",
        "stresse": "Stressé (faux amis)",
    }
    lignes = [
        [noms.get(p, p)] + [str(d.get(m, {}).get("par_profil", {}).get(p, "-")) for m in ("haiku", "sonnet")]
        for p in profils
    ]
    tableau(doc, ["Profil d'élève", "Haiku 4.5", "Sonnet 5"], lignes)


CORRECTIONS_R1 = [
    "Sécurité : numéros d'aide exacts avec leurs horaires réels (le 3018 est ouvert de 9 h à 23 h, pas 24 h/24), "
    "conduite claire face à un adulte rencontré en ligne (ne pas aller au rendez-vous, garder les messages, 119), "
    "ne pas reprendre l'exercice quand l'élève minimise une détresse, ne pas promettre le secret.",
    "Expressions de tous les jours (« ce contrôle va me tuer ») : ni alerte ni sermon.",
    "Réponse partielle : la forme conjuguée, le mot attendu ou le calcul avec les nombres de l'énoncé comptent "
    "comme la réponse.",
    "Rédaction : aucune phrase, aucun début de phrase à recopier.",
    "Vérifier avant de juger : refaire le calcul avant de dire « presque » ou « non ».",
    "Format : 2 à 5 phrases, une seule question par message ; consignes pour l'élève décrocheur et pour les "
    "remarques du profil (dys).",
    "Modes quiz et fiche : l'élève écrit sa fiche, le quiz ne change pas de sujet.",
    "Suivi et rapport : distinguer ce que l'élève a trouvé de ce que Jules a fourni ; ne pas dire « maîtrise ».",
    "Détection de la notion : relit les messages précédents, un « jsp » ne consomme plus d'essai.",
    "Coût : date à l'heure près et liste des notions dans un ordre fixe, pour que le cache du fournisseur s'applique ; "
    "cache activé dans le moteur anthropic.",
    "Garde-fou : le signe moins typographique (−3) n'échappe plus au détecteur.",
    "Contenu : fiche « aires urbaines » remise sur le zonage INSEE 2020 (aire d'attraction, couronne à 15 %, 93 %).",
]

CORRECTIONS_R2 = [
    "Garde-fou étendu : il relit aussi les messages libres tapés dans le panneau de Jules pendant une leçon, "
    "et plus seulement la réaction à une tentative (fuites « 12 × 12 = 144 » mesurées au tour 2).",
    "Correction automatique écrite dans la conversation : le suivi et le rapport voient qu'un exercice a été réussi.",
    "Formes irrégulières (went, took, bought) : les faire proposer par l'élève avant de les donner.",
    "Suivi : notion laissée vide tant que l'exercice n'est pas connu ; rapport accordé au genre de l'élève.",
]


def main() -> None:
    sortie = Path(sys.argv[1])
    chiffres = json.loads((ICI / "chiffres_campagne.json").read_text(encoding="utf-8"))
    calib_f = ICI.parents[2] / "jules-eval-eleves" / "evaluation" / "eleves" / "runs" / "r1" / "calibrage_manuel.json"
    calib = json.loads(calib_f.read_text(encoding="utf-8")) if calib_f.is_file() else {}
    concl_f = ICI / "conclusions.json"
    concl = json.loads(concl_f.read_text(encoding="utf-8")) if concl_f.is_file() else {}

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    t = doc.add_heading("Jules : tests avec des élèves simulés", level=0)
    for r in t.runs:
        r.font.color.rgb = BLEU
    para(
        doc,
        "Mesures d'efficacité et d'efficience, Haiku 4.5 et Sonnet 5. Campagne du 25 septembre 2026.",
        italique=True,
    )

    titre(doc, "L'essentiel")
    puces(doc, concl.get("essentiel", ["(conclusions à compléter)"]))

    titre(doc, "Comment on a testé")
    puces(
        doc,
        [
            "35 scénarios, 9 profils d'élèves (appliqué, « donne-moi la réponse », prétextes, injection, "
            "décrocheur, dys, brillant, distrait, détresse), 11 matières. Les 12 scénarios critiques (fuite, "
            "détresse) sont joués 3 fois : 59 conversations par modèle et par tour.",
            "Quatre surfaces : aide aux devoirs, leçon en blocs (mode cours), modes quiz, fiche et réexplique, "
            "épreuve sans aide. Tout le harnais tourne pour de vrai : détection de la notion, suivi, vigilance, "
            "rapport du soir.",
            "L'élève simulé est toujours joué par Sonnet 5 : seul Jules change d'un essai à l'autre.",
            "Correction par le code (fuite de la valeur attendue, longueur, questions) et par un juge Opus qui "
            "lit la conversation entière avec la réponse attendue. Opus sert seulement d'instrument de mesure, "
            "jamais de Jules.",
            "Trois tours : tour 1 sur le code de départ, corrections, tour 2, corrections, tour 3. Mêmes "
            "scénarios à chaque tour.",
            "Coût calculé au tarif API public (Haiku 4.5 : 1 $ / 5 $ par million de jetons ; Sonnet 5 : 2 $ / 10 $), "
            "sans cache. La latence est celle du CLI, pas celle de l'API.",
        ],
    )

    titre(doc, "Résultats, tour par tour")
    for modele, nom in (("haiku", "Haiku 4.5"), ("sonnet", "Sonnet 5")):
        titre(doc, nom, 2)
        table_tours(doc, chiffres, modele)
    dernier = max((t for t in chiffres if chiffres[t]), default="r1")
    titre(doc, f"Note du juge par profil d'élève (tour {dernier[1]})", 2)
    table_profils(doc, chiffres, dernier)

    titre(doc, "Ce qu'on a corrigé")
    titre(doc, "Après le tour 1", 2)
    puces(doc, CORRECTIONS_R1)
    titre(doc, "Après le tour 2", 2)
    puces(doc, CORRECTIONS_R2)

    titre(doc, "Fiabilité de la mesure")
    if calib:
        para(doc, calib.get("methode", ""))
        puces(doc, [f"{k.replace('_', ' ')} : {v}" for k, v in (calib.get("bilan") or {}).items()])
    puces(doc, concl.get("limites", []))

    titre(doc, "Efficience et coût")
    puces(doc, concl.get("cout", []))

    titre(doc, "Ce qui reste à faire")
    puces(doc, concl.get("reste", []))

    titre(doc, "Extraits de conversations")
    for ex in concl.get("extraits", []):
        para(doc, ex.get("titre", ""), gras=True)
        for ligne in ex.get("lignes", []):
            para(doc, ligne)
        if ex.get("commentaire"):
            p = doc.add_paragraph()
            r = p.add_run(ex["commentaire"])
            r.italic = True
            r.font.color.rgb = ROUGE

    sortie.parent.mkdir(parents=True, exist_ok=True)
    doc.save(sortie)
    print(sortie)


if __name__ == "__main__":
    main()
